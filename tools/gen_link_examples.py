#!/usr/bin/env python3
"""Regenerate tests/link_v2_examples.inc from the Examples sheet of the protocol workbook (the approved draft).

    python3 tools/gen_link_examples.py <path to slave_link_protocol_v2_draft.xlsx>            (writes the file)
    python3 tools/gen_link_examples.py <path to the workbook> --check                         (exit 1 if the file is out of date)

The test (tests/test_link_v2.c) parses every one of these frames and rebuilds it byte for byte, so the library cannot drift from what was approved."""
import os, sys
from openpyxl import load_workbook

root = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
target = os.path.join(root, "tests", "link_v2_examples.inc")
args = [a for a in sys.argv[1:] if not a.startswith("--")]
if len(args) != 1:
    sys.exit(__doc__)
ws = load_workbook(args[0], data_only=True)["Examples"]
rows = [(ws.cell(r, 1).value, ws.cell(r, 3).value) for r in range(5, ws.max_row + 1)
        if ws.cell(r, 3).value and str(ws.cell(r, 3).value).startswith("$")]
out = ("/* The example frames of the Examples sheet of slave_link_protocol_v2_draft.xlsx (the approved draft), copied by script. "
       "Do not edit by hand: change the workbook, then regenerate. */\n"
       "static const struct { const char *name; const char *frame; } EXAMPLES[] =\n{\n")
for name, frame in rows:
    out += '    { "%s", "%s" },\n' % (name, frame)
out += "};\n#define N_EXAMPLES ((int)(sizeof(EXAMPLES) / sizeof(EXAMPLES[0])))\n"
if "--check" in sys.argv:
    cur = open(target).read() if os.path.exists(target) else ""
    print("OK: %d examples, the file is up to date" % len(rows) if cur == out else "OUT OF DATE: regenerate tests/link_v2_examples.inc")
    sys.exit(0 if cur == out else 1)
open(target, "w").write(out)
print("wrote %s (%d examples)" % (target, len(rows)))
